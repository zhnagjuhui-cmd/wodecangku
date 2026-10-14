#!/usr/bin/env python3
"""Produce/reopen Chinese metadata documents only after evidence checks pass."""
import argparse
import hashlib
import json
from pathlib import Path
import zipfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

def main():
    p=argparse.ArgumentParser(); p.add_argument('--records',type=Path,required=True); p.add_argument('--verification',type=Path,required=True); p.add_argument('--out',type=Path,required=True); a=p.parse_args()
    verification=json.loads(a.verification.read_text()); records=json.loads(a.records.read_text())
    digest=hashlib.sha256(a.records.read_bytes()).hexdigest()
    if verification['status']!='PASS' or verification['records_sha256']!=digest:
        raise SystemExit('BLOCKED: evidence checks failed or records changed after verification')
    a.out.mkdir(parents=True,exist_ok=False)
    title='五项 ILD 研究文献元数据与证据入库报告'
    caution='仅核验原始研究元数据与摘要。未下载受限全文、未完成完整偏倚风险或 GRADE 评级；本文不构成治疗建议。两个数据库索引同一论文，不是独立临床证据。'
    doc=Document(); style=doc.styles['Normal']; style.font.name='Noto Sans CJK SC'; style.font.size=Pt(10)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    doc.add_heading(title,0); doc.add_paragraph(caution)
    lines=['# '+title,'',caution,'',f'记录哈希：`{digest}`','', '## 文献目录','', '| 研究 | 年份 | PMID | DOI | 核验 |', '| --- | --- | --- | --- | --- |']
    table=doc.add_table(rows=1,cols=4); table.style='Table Grid'
    for c,v in zip(table.rows[0].cells,['研究','年份','PMID','DOI']):c.text=v
    by_trial={r['trial']:r for r in verification['results']}
    for r in records:
        lines.append(f"| {r['trial']} | {r['year']} | [{r['pmid']}]({r['pubmed_url']}) | {r['doi']} | {by_trial[r['trial']]['status']} |")
        for c,v in zip(table.add_row().cells,[r['trial'],r['year'],r['pmid'],r['doi']]): c.text=v or '未核实'
    for r in records:
        first=r['authors'][0]; author=first['collective'] or ' '.join(filter(None,[first['last'],first['initials']]))
        details=[f"题名：{r['title']}",f"作者：{author} 等；期刊：{r['journal']}；{r['year']}，{r['volume']}({r['issue']})，{r['pages'] or '页码未核实'}。",
                 f"PMID：{r['pmid']}；DOI：{r['doi']}。",f"原始链接：{r['pubmed_url']}；{r['doi_url']}",
                 f"数据库研究类型：{'；'.join(r['publication_types'])}。本次仅原始论文摘要，非新闻公告或会议摘要。",
                 f"查询式：{r['query']}；候选数：{r['count']}；按原始题名唯一匹配。",
                 f"来源位置：{r['source_xpath']}；源响应：{r['pubmed_trace']['path']}；SHA-256：{r['pubmed_trace']['sha256']}；UTC 获取时间：{r['pubmed_trace']['utc']}。",
                 f"Europe PMC：同一 PMID/DOI/题名/年份及摘要数字集合核对通过（{by_trial[r['trial']]['numerical_token_count']} 个数字 token）；不代表统计学复算。",
                 '样本量、完整纳排标准、主要终点模型和临床结论尚未在本元数据测试中进行专业提取，不从摘要关键词推断填写。']
        doc.add_heading(r['trial'],1); lines.extend(['','## '+r['trial'],''])
        for s in details:doc.add_paragraph(s); lines.extend([s,''])
    if any(r['trial']=='U-ILD' for r in records):
        s='U-ILD 在本次测试中特指吡非尼酮用于未分类进展性纤维化 ILD 的 phase 2 原始研究，避免与其他未分类 ILD 研究混淆。'
        doc.add_paragraph(s);lines.extend(['',s])
    md=a.out/'ILD_METADATA_EVIDENCE.md'; md.write_text('\n'.join(lines),encoding='utf-8')
    path=a.out/'ILD_METADATA_EVIDENCE.docx'; doc.save(path)
    reopened=Document(path); body='\n'.join(p.text for p in reopened.paragraphs)+'\n'+'\n'.join(c.text for t in reopened.tables for row in t.rows for c in row.cells)
    with zipfile.ZipFile(path) as z:assert z.testzip() is None
    assert all(r['pmid'] in body and r['doi'] in body and r['title'] in body for r in records)
    assert len(reopened.tables[0].rows)==6 and caution in body
    result=dict(status='PASS',format='editable DOCX + Markdown',studies=5,records_sha256=digest,docx_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),citation_reopen_check=True)
    (a.out/'document_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps(result,ensure_ascii=False,indent=2))
    return 0

if __name__=='__main__':raise SystemExit(main())
