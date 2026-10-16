# coding: utf-8
"""Offline regression fixtures. Synthetic images are not model-generation evidence."""
import hashlib
import json
import os
from pathlib import Path
import random
import subprocess
import sys
import tempfile
import unittest
from PIL import Image

ROOT=Path(__file__).resolve().parent
ENV={**os.environ,'PYTHONDONTWRITEBYTECODE':'1'}


def run(script,*args,code=0):
    proc=subprocess.run([sys.executable,str(ROOT/script),*map(str,args)],capture_output=True,text=True,env=ENV)
    if proc.returncode!=code:
        raise AssertionError(proc.stdout+proc.stderr)
    return proc


def data(path):
    return json.loads(Path(path).read_text())


def put(path,value):
    Path(path).write_text(json.dumps(value,ensure_ascii=False))


def adapter(command,run_dir,**kwargs):
    argv=[command,'--run-dir',str(run_dir)]
    for key,value in kwargs.items():
        argv.append('--'+key.replace('_','-'))
        if value is not True:
            argv.append(str(value))
    return run('image_gen.py',*argv)


class PipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory()
        cls.base=Path(cls.temp.name).resolve()
        cls.deck=cls.base/'fixture'
        cls.deck.mkdir()
        rng=random.Random(12)
        image=Image.new('RGB',(256,144))
        image.putdata([(rng.randrange(256),rng.randrange(256),rng.randrange(256)) for _ in range(256*144)])
        cls.source=cls.base/'synthetic.png'
        image.save(cls.source)
        cls.policy={'target_size':'256x144','fit':'verify'}
        sample=cls.base/'sample.json'
        put(sample,{'slide':1,'prompt':'Synthetic offline fixture, not generated','image_policy':cls.policy})
        cls.sample_run=cls.deck/'sample-run'
        cls.make_image_job(sample,cls.sample_run)
        adapter('sample',cls.sample_run,deck=cls.deck,test_fixture=True)
        receipt=data(cls.deck/'sample_receipt.json')
        cls.spec={'deck_name':'fixture','language':'Chinese','image_policy':cls.policy,
                  'sample_generation_method':receipt['sample_generation_method'],
                  'style':{'name':'test'},'slides':[{'number':1,'title':'样张','sample_approved':True},
                  {'number':2,'title':'测试','key_points':['测试备注'],'required_images':[{'path':str(cls.source),'role':'content reference'}]}]}
        put(cls.deck/'deck_spec.json',cls.spec)
        run('prepare_slide_prompts.py','--spec',cls.deck/'deck_spec.json','--out-dir',cls.deck,'--test-fixture')
        cls.page_run=cls.deck/'image_runs/slide_02'
        cls.make_image_job(cls.deck/'prompts/slide_02.json',cls.page_run)
        run('record_slide_dispatch.py',cls.deck,'--slide','slide_02','--agent-id','fixture-agent')
        cls.output=data(cls.page_run/'result.json')['selected_source']
        run('record_slide_result.py',cls.deck,'--slide','slide_02','--agent-id','fixture-agent',
            '--backend-used','codex-imagegen-batch / built-in image tool','--image-run',cls.page_run,
            '--selected-source',cls.output,'--qa-note','Offline synthetic fixture only')
        (cls.deck/'speech.md').write_text('## Slide 1: 样张\n第一张备注。\n\n---\n注意点：核对顺序。\n\n## Slide 2: 测试\n第二张备注。\n',encoding='utf-8')

    @classmethod
    def make_image_job(cls,job,run_dir):
        run('image_gen.py','prepare','--job',job,'--run-dir',run_dir)
        adapter('claim',run_dir)
        adapter('record',run_dir,image=cls.source)
        adapter('finish',run_dir)
        adapter('approve',run_dir,qa_note='Synthetic offline test; no model used')

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_01_local_references_and_separate_queues(self):
        job=data(self.deck/'prompts/slide_02.json')
        self.assertEqual(len(job['input_images']),2)
        self.assertEqual(job['input_images'][0]['input_role'],'style')
        native=data(self.page_run/'queue/queue.json')['jobs'][0]
        self.assertEqual([x['role'] for x in native['references']],['style','reference'])
        self.assertNotEqual(self.sample_run,self.page_run)

    def test_02_state_and_packet_agree(self):
        jobs=data(self.deck/'slide_jobs.json')
        self.assertEqual([s['status'] for s in jobs['slides']],['accepted','recorded'])
        receipt=data(self.page_run/'result.json')
        self.assertEqual(receipt['attempts'],1)
        self.assertEqual(receipt['validation']['actual_size'],[256,144])
        self.assertEqual(receipt['backend_family'],'built-in-image-tool')

    def test_03_test_sample_does_not_authorize_production(self):
        run('prepare_slide_prompts.py','--spec',self.deck/'deck_spec.json','--out-dir',self.deck,'--force',code=1)
        self.assertFalse(data(self.deck/'sample_receipt.json')['user_approved'])

    def test_04_resume_keeps_completed_results(self):
        run('image_gen.py','prepare','--job',self.deck/'prompts/slide_02.json','--run-dir',self.page_run)
        self.assertEqual(data(self.page_run/'queue/queue.json')['jobs'][0]['attempts'],1)
        adapter('claim',self.page_run)
        self.assertEqual(data(self.page_run/'queue/queue.json')['jobs'][0]['status'],'completed')

    def test_05_final_file_tamper_blocks_assembly(self):
        from preflight import validate_deck
        path=self.deck/'origin_image/slide_02.png'
        before=path.read_bytes()
        try:
            path.write_bytes(b'not an image')
            with self.assertRaises(ValueError):
                validate_deck(self.deck)
        finally:
            path.write_bytes(before)

    def test_06_incomplete_page_blocks_assembly(self):
        from preflight import validate_deck
        path=self.deck/'slide_jobs.json'
        before=path.read_bytes()
        try:
            jobs=data(path)
            jobs['slides'][1]['status']='dispatched'
            put(path,jobs)
            with self.assertRaises(ValueError):
                validate_deck(self.deck)
        finally:
            path.write_bytes(before)

    def test_07_missing_notes_block_assembly(self):
        from preflight import validate_deck
        path=self.deck/'speech.md'
        before=path.read_bytes()
        try:
            path.write_text('## Slide 1: 备注\n只有一页。')
            with self.assertRaises(ValueError):
                validate_deck(self.deck)
        finally:
            path.write_bytes(before)

    def test_08_urls_and_unapproved_postprocessing_are_rejected(self):
        for name,entry in [('url',{'input_images':[{'path':'https://example.com/a.png'}]}),
                           ('fit',{'image_policy':{'fit':'contain'}})]:
            path=self.base/(name+'.json')
            put(path,{'slide':3,'prompt':'test',**entry})
            run('image_gen.py','prepare','--job',path,'--run-dir',self.base/name,code=2)

    def test_09_uncertain_retry_needs_ack_and_quality_needs_policy(self):
        path=self.base/'retry.json'
        put(path,{'slide':3,'prompt':'test'})
        run_dir=self.base/'retry-run'
        run('image_gen.py','prepare','--job',path,'--run-dir',run_dir)
        adapter('claim',run_dir)
        run('image_gen.py','retry','--run-dir',run_dir,'--reason','uncertain',code=2)
        run('image_gen.py','retry','--run-dir',run_dir,'--reason','quality',code=2)
        adapter('retry',run_dir,reason='uncertain',ack_duplicate_risk=True)
        self.assertEqual(data(run_dir/'queue/queue.json')['jobs'][0]['status'],'pending')

    def test_10_legacy_commands_fail_explicitly(self):
        for name in ('generate','generate-batch','edit'):
            run('image_gen.py',name,code=2)

    def test_11_assembly_keeps_page_order_and_chinese_notes(self):
        from pptx import Presentation
        run('assemble_ppt.py',self.base,'fixture.pptx')
        p=Presentation(self.deck/'fixture.pptx')
        self.assertEqual(len(p.slides),2)
        self.assertAlmostEqual(p.slide_width/p.slide_height,16/9)
        self.assertIn('第一张备注',p.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn('第二张备注',p.slides[1].notes_slide.notes_text_frame.text)
        self.assertEqual([len(s.shapes) for s in p.slides],[1,1])

    def test_12_different_page_run_cannot_be_received(self):
        path=self.deck/'slide_jobs.json'
        before=path.read_bytes()
        try:
            jobs=data(path)
            jobs['slides'][1]['status']='dispatched'
            put(path,jobs)
            p=run('record_slide_result.py',self.deck,'--slide','slide_02','--agent-id','fixture-agent',
                '--backend-used','codex-imagegen-batch / built-in image tool','--image-run',self.sample_run,
                '--selected-source',self.output,'--qa-note','wrong page',code=1)
            self.assertIn('different slide',p.stderr)
        finally:
            path.write_bytes(before)

    def test_13_alternate_backend_rejected(self):
        path=self.deck/'slide_jobs.json'
        before=path.read_bytes()
        try:
            jobs=data(path)
            jobs['slides'][1]['status']='dispatched'
            put(path,jobs)
            p=run('record_slide_result.py',self.deck,'--slide','slide_02','--agent-id','fixture-agent',
                '--backend-used','external image provider','--image-run',self.page_run,
                '--selected-source',self.output,'--qa-note','wrong backend',code=1)
            self.assertIn('fixed Codex',p.stderr)
        finally:
            path.write_bytes(before)

    def test_14_authorized_fit_preserves_original(self):
        path=self.base/'fit-allowed.json'
        put(path,{'slide':4,'prompt':'fixture','image_policy':{'target_size':'512x288','fit':'contain','approved':True}})
        run_dir=self.base/'fit-allowed'
        self.make_image_job(path,run_dir)
        result=data(run_dir/'result.json')
        self.assertEqual(result['validation']['actual_size'],[512,288])
        self.assertTrue(result['validation']['upscaled'])
        self.assertEqual(Path(result['generated_original']).read_bytes(),self.source.read_bytes())

    def test_15_quality_retry_budget_is_bounded(self):
        path=self.base/'quality.json'
        put(path,{'slide':5,'prompt':'fixture','image_policy':{'target_size':'256x144','approved':True,'auto_quality_retry':True,'max_attempts':2}})
        run_dir=self.base/'quality-run'
        run('image_gen.py','prepare','--job',path,'--run-dir',run_dir)
        for attempt in (1,2):
            adapter('claim',run_dir)
            adapter('record',run_dir,image=self.source)
            adapter('finish',run_dir)
            if attempt==1:
                adapter('retry',run_dir,reason='quality')
        run('image_gen.py','retry','--run-dir',run_dir,'--reason','quality',code=2)
        self.assertEqual(data(run_dir/'queue/queue.json')['jobs'][0]['attempts'],2)

    def test_16_dispatched_state_cannot_be_reset(self):
        p=run('prepare_slide_prompts.py','--spec',self.deck/'deck_spec.json','--out-dir',self.deck,'--test-fixture','--force',code=1)
        self.assertIn('must not be reset',p.stderr)

    def receive_page(self,code=0):
        return run('record_slide_result.py',self.deck,'--slide','slide_02','--agent-id','fixture-agent',
                   '--backend-used','codex-imagegen-batch / built-in image tool','--image-run',self.page_run,
                   '--selected-source',self.output,'--qa-note','Offline fixture',code=code)

    def test_17_old_prompt_result_rejected_after_new_dispatch(self):
        path=self.deck/'slide_jobs.json'
        prompt=self.deck/'prompts/slide_02.json'
        state_before,prompt_before=path.read_bytes(),prompt.read_bytes()
        try:
            j=data(prompt)
            j['prompt']='Changed prompt after old queue was completed'
            put(prompt,j)
            jobs=data(path)
            jobs['slides'][1]['status']='dispatched'
            jobs['slides'][1]['dispatch']['prompt_sha256']=hashlib.sha256(prompt.read_bytes()).hexdigest()
            put(path,jobs)
            p=self.receive_page(code=1)
            self.assertIn('not prepared from',p.stderr)
        finally:
            path.write_bytes(state_before)
            prompt.write_bytes(prompt_before)

    def test_18_native_record_interruption_recovers_without_regeneration(self):
        path=self.page_run/'bridge.json'
        before=path.read_bytes()
        try:
            bridge=data(path)
            bridge.pop('generation')
            put(path,bridge)
            adapter('record',self.page_run,image=self.source)
            adapter('packet',self.page_run)
            self.assertTrue(data(path)['generation']['recovered'])
            self.assertEqual(data(self.page_run/'queue/queue.json')['jobs'][0]['attempts'],1)
        finally:
            path.write_bytes(before)

    def test_19_identical_copied_result_recovers_dispatch_state(self):
        path=self.deck/'slide_jobs.json'
        before=path.read_bytes()
        try:
            jobs=data(path)
            jobs['slides'][1]['status']='dispatched'
            put(path,jobs)
            self.receive_page()
            self.assertEqual(data(path)['slides'][1]['status'],'recorded')
        finally:
            path.write_bytes(before)

    def test_20_sample_copy_interruption_recovers_receipt(self):
        path=self.deck/'sample_receipt.json'
        before=path.read_bytes()
        try:
            path.unlink()
            adapter('sample',self.sample_run,deck=self.deck,test_fixture=True)
            self.assertTrue(data(path)['test_fixture'])
            self.assertFalse(data(path)['user_approved'])
        finally:
            path.write_bytes(before)

    def test_21_page_size_override_reaches_prompt_and_queue(self):
        deck=self.base/'size-override'
        deck.mkdir()
        adapter('sample',self.sample_run,deck=deck,test_fixture=True)
        spec={**self.spec,'deck_name':'size-override','slides':[dict(s) for s in self.spec['slides']]}
        spec['slides'][1]['image_policy']={'target_size':'512x288'}
        spec['sample_generation_method']=data(deck/'sample_receipt.json')['sample_generation_method']
        put(deck/'deck_spec.json',spec)
        run('prepare_slide_prompts.py','--spec',deck/'deck_spec.json','--out-dir',deck,'--test-fixture')
        prompt=data(deck/'prompts/slide_02.json')
        self.assertIn('"requested_size": "512x288"',prompt['prompt'])
        run('image_gen.py','prepare','--job',deck/'prompts/slide_02.json','--run-dir',deck/'image_runs/slide_02')
        self.assertEqual(data(deck/'image_runs/slide_02/queue/queue.json')['jobs'][0]['size'],'512x288')

    def test_22_blocked_page_resume_keeps_image_queue(self):
        path=self.deck/'slide_jobs.json'
        before=path.read_bytes()
        try:
            jobs=data(path)
            jobs['slides'][1]['status']='blocked'
            jobs['slides'][1]['blocker']={'reason':'interrupted'}
            put(path,jobs)
            run('resume_slide.py',self.deck,'--slide','slide_02','--agent-id','wrong','--reason','resume',code=1)
            run('resume_slide.py',self.deck,'--slide','slide_02','--agent-id','fixture-agent','--reason','resume')
            self.assertEqual(data(path)['slides'][1]['status'],'dispatched')
            self.assertEqual(len(data(path)['slides'][1]['recovery_history']),1)
            self.assertEqual(data(self.page_run/'queue/queue.json')['jobs'][0]['attempts'],1)
        finally:
            path.write_bytes(before)

    def test_23_existing_different_image_is_not_overwritten(self):
        path=self.deck/'slide_jobs.json'
        target=self.deck/'origin_image/slide_02.png'
        state_before,image_before=path.read_bytes(),target.read_bytes()
        try:
            jobs=data(path)
            jobs['slides'][1]['status']='dispatched'
            put(path,jobs)
            Image.new('RGB',(256,144),'blue').save(target)
            changed=target.read_bytes()
            self.receive_page(code=1)
            self.assertEqual(target.read_bytes(),changed)
            self.assertEqual(data(path)['slides'][1]['status'],'dispatched')
        finally:
            path.write_bytes(state_before)
            target.write_bytes(image_before)


if __name__=='__main__':
    unittest.main()
