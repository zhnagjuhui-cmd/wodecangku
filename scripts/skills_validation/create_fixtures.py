#!/usr/bin/env python3
"""Synthetic component fixtures; no clinical production deck or generated imagery."""
import json
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def main():
    root=Path(__file__).resolve().parents[2]; out=root/'outputs/skills_validation_20261014/fixtures'; out.mkdir(exist_ok=False)
    image=Image.new('RGB',(1600,900),'white'); draw=ImageDraw.Draw(image)
    font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',44)
    small=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',28)
    draw.text((100,100),'SENSCIS metadata fixture',font=font,fill='#8C0D27')
    draw.rectangle((100,250,1500,420),fill='#087B78');draw.text((140,280),'PMID 31112379 - abstract only',font=font,fill='white')
    draw.text((100,730),'Source: PubMed / DOI 10.1056/NEJMoa1903076',font=small,fill='#444444')
    image.save(out/'source.png')
    p=Presentation();p.slide_width=Inches(16);p.slide_height=Inches(9);s=p.slides.add_slide(p.slide_layouts[6])
    def box(name,x,y,w,h,value,size,color,fill=None):
        b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));b.name=name;b.text=value
        if fill:b.fill.solid();b.fill.fore_color.rgb=RGBColor.from_string(fill)
        tf=b.text_frame;tf.margin_left=Inches(.1);tf.margin_top=Inches(.05)
        for para in tf.paragraphs:
            for run in para.runs:run.font.size=Pt(size);run.font.name='DejaVu Sans';run.font.color.rgb=RGBColor.from_string(color)
    box('title',1,1,14,1,'SENSCIS metadata fixture',32,'8C0D27')
    box('conclusion',1,2.5,14,1.7,'PMID 31112379 - abstract only',32,'FFFFFF','087B78')
    box('reference',1,7.3,14,1,'Source: PubMed / DOI 10.1056/NEJMoa1903076',20,'444444')
    s.notes_slide.notes_text_frame.text='本页为运行验证测试夹具，不是临床课件。仅包含真实文献标识，不包含试验疗效结论。结论与引文为独立可编辑文本对象。'
    p.save(out/'reconstructed.pptx')
    (out/'bounds.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="900" viewBox="0 0 1600 900"><text x="100" y="150" font-size="44">SENSCIS metadata fixture</text><text x="100" y="760" font-size="28">PMID 31112379</text></svg>')
    kai=out/'kaiyuan_fixture';(kai/'origin_image').mkdir(parents=True);image.save(kai/'origin_image/slide_01.png')
    (kai/'speech.md').write_text('## Slide 1: metadata fixture\n\n独立组装器测试夹具。该页为图片式 PPT，不声称原生可编辑，也未调用生图服务。\n')
    template=(root/'.agents/skills/guizang-ppt-skill/assets/template-swiss.html').read_text()
    # Existing seed content is exercised as a fixture, not delivered as a deck.
    template=template.replace('[必填] 替换为 PPT 标题 · Deck Title','Guizang runtime fixture')
    (out/'guizang_seed.html').write_text(template)
    brief=[dict(role='cover',priority='title',content=dict(presentation=dict(title='文献入库验证',titleShort='文献验证',summary='仅元数据测试，无临床推荐',summaryShort='元数据测试',takeaway='保留原始来源',items=[]),meta=dict(brand='ILD',pageLabel='2026',panelTitle='验证')))]
    (out/'dashi_brief.json').write_text(json.dumps(brief,ensure_ascii=False,indent=2))
    print('PASS: synthetic fixtures created in',out)

if __name__=='__main__':main()
