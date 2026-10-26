"""Render an existing PPTX without changing it; require a new output directory."""
import argparse, pathlib, shutil, subprocess, tempfile, json, hashlib
import fitz
from pptx import Presentation
p=argparse.ArgumentParser();p.add_argument('pptx',type=pathlib.Path);p.add_argument('output',type=pathlib.Path);p.add_argument('--citation',default='');a=p.parse_args()
source=a.pptx.resolve();out=a.output.resolve();out.mkdir(parents=True,exist_ok=False)
exe=shutil.which('soffice') or shutil.which('libreoffice')
if not exe:raise SystemExit('No soffice/libreoffice executable')
with tempfile.TemporaryDirectory(prefix='ild-lo-') as tmp:
 profile=pathlib.Path(tmp)/'profile'
 run=subprocess.run([exe,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(out),str(source)],capture_output=True,text=True,timeout=120)
 (out/'conversion.log').write_text(run.stdout+run.stderr)
 if run.returncode:raise SystemExit(run.returncode)
pdf=out/(source.stem+'.pdf');doc=fitz.open(pdf);deck=Presentation(source)
assert len(doc)==len(deck.slides)
report={'renderer':subprocess.check_output([exe,'--version'],text=True).strip(),'pptx_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'pages':len(doc),'slides':[],'visual_review':'pending; inspect each PNG'}
for i,page in enumerate(doc,1):
 page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(out/f'slide-{i}.png')
 text=page.get_text();(out/f'slide-{i}.txt').write_text(text)
 note=deck.slides[i-1].notes_slide.notes_text_frame.text
 assert len(note)>0
 if a.citation:assert a.citation in text
 report['slides'].append({'page':i,'notes_chars':len(note),'citation_present':a.citation in text if a.citation else None,'fonts':sorted(set(x[3] for x in page.get_fonts()))})
(out/'render_check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({'pdf':str(pdf),'pages':len(doc)},ensure_ascii=False))
